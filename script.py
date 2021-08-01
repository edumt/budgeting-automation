import os
import datetime
import re
from openpyxl import load_workbook
from docx import Document

# from docx2pdf import convert #TODO: implement export docx as pdf

SHEET_VERSION = "13-JULHO-2021"
CONSUMPTION_TABLE = "HCONSUMO"
INVERTER_PANEL = "Preço SFCR-GROWATT PHONO 450Wp"
NEAREST_PANELS = "=+VLOOKUP(MIN('GROWATT-PHONO 450Wp'!D:D),'GROWATT-PHONO 450Wp'!D:E,2,0)"

# sheet cells constants
CONSUMPTION = "D18"
NAME = "C4"
PRICE_ADJUSTMENT = "H44"  # TODO: dinamically adjust price
OPTION_1 = "D11"
OPTION_2 = "F11"
OPTION_3 = "H11"
# QUOTE_AREA = "B3:H32"

statesDict = {
    "ES": "ESPÍRITO SANTO",
    "RJ": "RIO DE JANEIRO",
    "BA": "BAHIA",
    "MG": "MINAS GERAIS",
}


def getFormulaSheet():
    return load_workbook("../excel_template.xlsx")


def getDocxTemplate():
    return Document("../word_template.docx")


def getCurrentDate():
    return datetime.date.today().strftime("%d/%m/%Y")


def getReferenceNumber():
    def filterTempFiles(name):
        if name[0] == "9":
            return True
        else:
            return False

    filtered_dir_list = filter(filterTempFiles, os.listdir("../DOCS"))
    latest_file = sorted(filtered_dir_list, reverse=True)[0]
    reference = re.search("(?<=999-PTC-70-)\d*(?=-2)", latest_file)
    return int(reference.group(0)) + 1


class Client:
    def __init__(self, name, consumption, state, city=""):
        self.name = name.upper()
        self.consumption = consumption  # Average monthly energy consumption [kWh/month]
        self.state = state.upper()  # State initials, e.g. ES
        self.city = city.upper()  # (District-)City, e.g. PRAIA DE ITAPARICA-VILA VELHA
        self.reference = getReferenceNumber()

    def populateSheet(self):
        self.sheet[CONSUMPTION_TABLE][CONSUMPTION].value = self.consumption
        self.sheet[CONSUMPTION_TABLE][NAME].value = f"CLIENTE: {self.name}"
        self.setPanelsQuantity()

    def setPanelsQuantity(self):
        if self.consumption <= 550:
            self.sheet[INVERTER_PANEL][OPTION_2].value = NEAREST_PANELS + "+1"
            self.sheet[INVERTER_PANEL][OPTION_3].value = NEAREST_PANELS + "+2"

        elif self.consumption <= 750:
            self.sheet[INVERTER_PANEL][OPTION_1].value = NEAREST_PANELS + "-1"
            self.sheet[INVERTER_PANEL][OPTION_2].value = NEAREST_PANELS
            self.sheet[INVERTER_PANEL][OPTION_3].value = NEAREST_PANELS + "+1"

        elif self.consumption <= 1300:
            self.sheet[INVERTER_PANEL][OPTION_1].value = NEAREST_PANELS + "-2"
            self.sheet[INVERTER_PANEL][OPTION_2].value = NEAREST_PANELS + "-1"
            self.sheet[INVERTER_PANEL][OPTION_3].value = NEAREST_PANELS

        else:
            self.sheet[INVERTER_PANEL][OPTION_1].value = NEAREST_PANELS + "-4"
            self.sheet[INVERTER_PANEL][OPTION_2].value = NEAREST_PANELS + "-2"
            self.sheet[INVERTER_PANEL][OPTION_3].value = NEAREST_PANELS

    def saveSheet(self):
        # try:
        #    # f"../{statesDict[self.state]}/{self.name}""
        #    os.mkdir(f"../SHEETS/{self.name}")
        # except OSError as error:
        #   if error.errno == 17:
        #      # If directory already exists do nothing
        #     pass
        # else:
        #   # prob just raise exception again instead of printing
        #   raise
        #    #print(error)
        # self.sheet.save(f"../SHEETS/{self.name}/GERADORES-{self.name}-ALDO-{SHEET_VERSION}.xlsx")
        self.sheet.save("../test.xlsx")

    def generateSheet(self):
        self.sheet = getFormulaSheet()
        self.populateSheet()
        self.saveSheet()
        print("Success! Sheet generated.")

    def docxSearchAndReplace(self, oldText, newText):
        for p in self.docx.paragraphs:
            if oldText in p.text:
                for r in p.runs:
                    if oldText == r.text:
                        r.text = newText

    def populateDocx(self):
        self.docxSearchAndReplace("fullName", self.name)
        self.docxSearchAndReplace("reference", f"{self.reference}")
        self.docxSearchAndReplace("date", getCurrentDate())
        if self.city == "":
            self.docxSearchAndReplace("location", statesDict[self.state])
        else:
            self.docxSearchAndReplace("location", f"{self.city}/{self.state}")
        self.docxSearchAndReplace("consumption", f"{self.consumption:,}".replace(",", "."))

    def saveDocx(self):
        # self.docx.save(f"../DOCS/999-PTC-70-{self.reference}-21 R0 (SFCR-{self.name}).docx")
        self.docx.save("../test.docx")

    def generateQuote(self):
        self.docx = getDocxTemplate()
        self.populateDocx()
        self.saveDocx()
        print("Success! Quote generated.")


def main():
    test_client = Client(
        "Eduardo Moura Tavares",
        500,
        "ES",
        "Praia de Itaparica-Vila Velha",
    )
    # test_client = Client('Eduardo Moura Tavares', 5000, 'mg')
    test_client.generateSheet()
    test_client.generateQuote()
    os.startfile("F:/Google Drive/Projetos/test.docx")  # for testing purposes
    os.startfile("F:/Google Drive/Projetos/test.xlsx")  # for testing purposes


if __name__ == "__main__":
    main()
"""     
    TO DO:  GUI
            
    MAYBE TO DO:
                e-mail integration
                whatsapp integration
"""
